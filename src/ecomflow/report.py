"""报表层：自动生成多 Sheet Excel 日报 + 可直接进邮件的 HTML 摘要。"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import metrics
from .config import Config, get_config
from .logging_setup import get_logger

log = get_logger(__name__)

NAVY = "1F4E79"
ORANGE = "C55A11"
LIGHT = "DDEBF7"
GREY = "F2F2F2"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor=LIGHT)
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=15, color=NAVY)
NOTE_FONT = Font(size=9, color="7F7F7F")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LEVEL_FILL = {"high": PatternFill("solid", fgColor="F8CBAD"),
              "medium": PatternFill("solid", fgColor="FFE699"),
              "low": PatternFill("solid", fgColor="E2EFDA")}
STATUS_FILL = {"断货": "F8CBAD", "紧急补货": "FFE699", "低于安全库存": "FFF2CC",
               "滞销": "DDEBF7", "健康": "E2EFDA"}

CN_COLUMNS = {
    "date": "日期", "period": "周期", "platform": "平台", "category": "品类",
    "channel": "渠道", "sku_id": "SKU编码", "sku_name": "商品名称", "warehouse": "仓库",
    "gmv": "GMV(元)", "net_gmv": "净销售额(元)", "refund_amount": "退款金额(元)",
    "orders": "订单数", "buyers": "买家数", "units": "销量", "cost": "成本(元)",
    "margin": "毛利(元)", "aov": "客单价(元)", "margin_rate": "毛利率",
    "refund_rate": "退款率", "gmv_share": "GMV占比", "cum_share": "累计占比",
    "abc": "ABC", "impressions": "曝光量", "clicks": "点击量", "spend": "花费(元)",
    "attr_orders": "归因订单", "attr_gmv": "归因GMV(元)", "ctr": "点击率",
    "cpc": "平均点击成本", "cvr": "转化率", "roi": "ROI", "gmv_change": "环比",
    "on_hand": "现有库存", "in_transit": "在途库存", "safety_stock": "安全库存",
    "avg_daily_sales": "日均销量", "cover_days": "可售天数", "slow_days": "无动销天数",
    "status": "状态", "suggest_qty": "建议补货", "suggestion": "处理建议",
    "campaign_id": "活动ID", "campaign_name": "活动名称", "campaign_type": "类型",
    "start_date": "开始日期", "end_date": "结束日期", "days": "天数", "budget": "预算(元)",
    "discount_rate": "折扣率", "base_gmv": "基准期GMV(元)", "lift_pct": "提升幅度",
    "incremental_gmv": "增量GMV(元)", "conclusion": "结论", "rule_id": "规则",
    "rule_name": "规则名称", "level": "等级", "scope": "范围", "message": "异常描述",
    "detail_key": "定位", "metric": "指标值",
    "first_seen": "首次发现", "last_seen": "最近发现", "occurrences": "累计次数",
    "owner": "处理人", "issue_id": "工单号", "note": "备注",
    "table": "数据表", "rule": "清洗规则", "hit_rows": "命中行数",
    "source": "数据源", "type": "类型", "rows": "行数", "columns": "列数",
    "encoding": "编码", "elapsed_ms": "耗时(ms)", "message_y": "采集信息",
}


def _cn(name: str) -> str:
    return CN_COLUMNS.get(name, name)


def _fmt_value(v: Any) -> Any:
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, (date,)):
        return v.isoformat()
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating,)):
        v = float(v)
    if isinstance(v, (np.bool_,)):
        return bool(v)
    if isinstance(v, float) and pd.isna(v):
        return None
    if v is pd.NaT:
        return None
    return v


def _text_width(value: Any) -> int:
    text = "" if value is None else str(value)
    return sum(2 if ord(ch) > 127 else 1 for ch in text)


def _write_table(ws: Worksheet, df: pd.DataFrame, title: str, subtitle: str = "",
                 start_row: int = 1, formats: dict[str, str] | None = None,
                 color_col: str | None = None, colors: dict[str, str] | None = None,
                 max_width: int = 34) -> int:
    """在工作表指定位置写一张带样式的表，返回下一张表的起始行。"""
    ws.cell(row=start_row, column=1, value=title).font = TITLE_FONT
    if subtitle:
        cell = ws.cell(row=start_row + 1, column=1, value=subtitle)
        cell.font = NOTE_FONT
    header_row = start_row + 2
    cols = list(df.columns)
    for j, col in enumerate(cols, start=1):
        cell = ws.cell(row=header_row, column=j, value=_cn(col))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    for i, row in enumerate(df.itertuples(index=False), start=header_row + 1):
        for j, (col, value) in enumerate(zip(cols, row), start=1):
            cell = ws.cell(row=i, column=j, value=_fmt_value(value))
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center",
                                       horizontal="left" if isinstance(value, str) else "right")
            fmt = (formats or {}).get(col)
            if fmt:
                cell.number_format = fmt
            if color_col == col and colors:
                fill = colors.get(str(value))
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
    # 列宽
    for j, col in enumerate(cols, start=1):
        width = max([_text_width(_cn(col))] + [_text_width(_fmt_value(v)) for v in df[col].head(50)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 8), max_width)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if len(df):
        ws.auto_filter.ref = (f"A{header_row}:{get_column_letter(len(cols))}"
                              f"{header_row + len(df)}")
    ws.row_dimensions[header_row].height = 22
    return header_row + len(df) + 3


def _kpi_block(ws: Worksheet, kpi: dict, meta: dict, quality_score: float,
               issue_counts: dict) -> None:
    ws["A1"] = "电商运营数据日报"
    ws["A1"].font = Font(bold=True, size=20, color=NAVY)
    ws["A2"] = (f"数据区间 {meta.get('date_start', '-')} ~ {meta.get('date_end', '-')}   |   "
                f"生成时间 {meta.get('finished_at', '-')}   |   批次 {meta.get('run_id', '-')}")
    ws["A2"].font = NOTE_FONT

    cards = [
        ("近30天 GMV(元)", kpi.get("gmv", 0), "gmv_change", "#,##0"),
        ("近30天 净销售额(元)", kpi.get("net_gmv", 0), "net_gmv_change", "#,##0"),
        ("近30天 订单数", kpi.get("orders", 0), "orders_change", "#,##0"),
        ("客单价(元)", kpi.get("aov", 0), "aov_change", "#,##0.00"),
        ("毛利率", kpi.get("margin_rate", 0), "margin_rate_change", "0.0%"),
        ("退款率", kpi.get("refund_rate", 0), "refund_rate_change", "0.00%"),
        ("数据质量分", quality_score, None, "0.0"),
        ("待处理异常", issue_counts.get("total", 0), None, "0"),
    ]
    for idx, (label, value, change_key, fmt) in enumerate(cards):
        col = 1 + (idx % 4) * 3
        row = 4 + (idx // 4) * 3
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        label_cell = ws.cell(row=row, column=col, value=label)
        label_cell.font = Font(size=9, color="595959")
        label_cell.fill = SUB_FILL
        label_cell.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        value_cell = ws.cell(row=row + 1, column=col, value=float(value or 0))
        value_cell.font = Font(bold=True, size=16, color=NAVY)
        value_cell.number_format = fmt
        value_cell.alignment = Alignment(horizontal="center")
        if change_key:
            change = kpi.get(change_key)
            text = "—" if change is None or pd.isna(change) else f"环比 {change:+.1%}"
            chg = ws.cell(row=row + 2, column=col, value=text)
            chg.font = Font(size=9, color="C00000" if (change or 0) < 0 else "548235")
            chg.alignment = Alignment(horizontal="center")

    ws["A11"] = ("口径：GMV=支付成功+已退款订单实付；净销售额=支付成功实付-退款；"
                 "毛利=实付-成本×数量；数据质量分=100-按等级扣分。")
    ws["A11"].font = NOTE_FONT
    for c in range(1, 13):
        ws.column_dimensions[get_column_letter(c)].width = 16


def build_excel(path: Path, tables: dict[str, pd.DataFrame], issues: pd.DataFrame,
                daily: pd.DataFrame, kpi: dict, clean_log: pd.DataFrame,
                ingest_log: pd.DataFrame, meta: dict, cfg: Config) -> Path:
    orders = tables.get("orders", pd.DataFrame())
    inventory = tables.get("inventory", pd.DataFrame())
    traffic = tables.get("traffic", pd.DataFrame())
    campaigns = tables.get("campaigns", pd.DataFrame())
    top_n = int(cfg.get("report.top_n", 20))
    quality = metrics.data_quality_score(issues)
    counts = {"total": len(issues)}

    wb = Workbook()
    ws = wb.active
    ws.title = "经营概览"
    _kpi_block(ws, kpi, meta, quality, counts)

    # 日销售趋势 + 折线图
    trend = daily.copy()
    trend["date"] = pd.to_datetime(trend["date"]).dt.strftime("%Y-%m-%d")
    trend = trend[["date", "gmv", "net_gmv", "orders", "aov", "margin_rate", "refund_rate"]]
    ws2 = wb.create_sheet("日销售趋势")
    _write_table(ws2, trend, "日销售趋势",
                 "GMV 含退款单，净销售额为支付成功口径；毛利率按实付口径计算",
                 formats={"gmv": "#,##0", "net_gmv": "#,##0", "orders": "#,##0",
                          "aov": "#,##0.00", "margin_rate": "0.0%", "refund_rate": "0.00%"})
    if len(trend) > 1:
        chart = LineChart()
        chart.title = "日 GMV 趋势"
        chart.style = 12
        chart.height, chart.width = 9, 26
        chart.y_axis.title = "GMV(元)"
        data = Reference(ws2, min_col=2, min_row=3, max_row=3 + len(trend))
        cats = Reference(ws2, min_col=1, min_row=4, max_row=3 + len(trend))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws2.add_chart(chart, f"A{8 + len(trend)}")

    # 平台 / 品类
    plat = metrics.dimension_summary(orders, "platform", days=30)
    ws3 = wb.create_sheet("平台表现")
    _write_table(ws3, plat, "平台表现（近 30 天）", "按 GMV 降序，退款率=退款金额/GMV",
                 formats={"gmv": "#,##0", "net_gmv": "#,##0", "refund_amount": "#,##0",
                          "orders": "#,##0", "buyers": "#,##0", "units": "#,##0",
                          "margin": "#,##0", "margin_rate": "0.0%",
                          "refund_rate": "0.00%", "gmv_share": "0.0%"})
    if len(plat) > 1:
        chart = BarChart()
        chart.title = "各平台 GMV"
        chart.height, chart.width = 7, 18
        data = Reference(ws3, min_col=2, min_row=3, max_row=3 + len(plat))
        cats = Reference(ws3, min_col=1, min_row=4, max_row=3 + len(plat))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws3.add_chart(chart, f"A{8 + len(plat)}")

    cat = metrics.dimension_summary(orders, "category", days=30)
    _write_table(wb.create_sheet("品类表现"), cat, "品类表现（近 30 天）",
                 "关注毛利率与 GMV 占比的匹配度，低毛利高占比品类优先优化",
                 formats={"gmv": "#,##0", "net_gmv": "#,##0", "orders": "#,##0",
                          "margin": "#,##0", "margin_rate": "0.0%",
                          "refund_rate": "0.00%", "gmv_share": "0.0%"})

    sku = metrics.sku_ranking(orders, n=top_n, days=30)
    _write_table(wb.create_sheet("SKU排行"), sku, f"SKU 销售排行 TOP{top_n}（近 30 天）",
                 "A 类=累计贡献 70% GMV，B 类=70%~90%，C 类=后 10%",
                 formats={"gmv": "#,##0", "units": "#,##0", "orders": "#,##0",
                          "margin": "#,##0", "margin_rate": "0.0%",
                          "gmv_share": "0.0%", "cum_share": "0.0%"})

    if not inventory.empty:
        alert = metrics.inventory_alert(inventory, orders, cfg)
        ws5 = wb.create_sheet("库存预警")
        _write_table(ws5, alert, "库存健康度与补货建议",
                     "可售天数=(现有+在途)/近30天日均销量；滞销=超过 45 天无动销",
                     formats={"avg_daily_sales": "0.00", "cover_days": "0.0",
                              "suggest_qty": "#,##0"},
                     color_col="status", colors=STATUS_FILL, max_width=40)

    if not campaigns.empty:
        eff = metrics.campaign_effect(orders, campaigns, cfg)
        _write_table(wb.create_sheet("活动效果"), eff, "活动效果评估",
                     "基准期=活动开始前等长周期；提升幅度=(活动期GMV-基准期GMV)/基准期GMV",
                     formats={"gmv": "#,##0", "orders": "#,##0", "aov": "#,##0",
                              "budget": "#,##0", "roi": "0.00", "base_gmv": "#,##0",
                              "lift_pct": "0.0%", "incremental_gmv": "#,##0",
                              "discount_rate": "0.0%"},
                     max_width=42)

    if not traffic.empty:
        tr = metrics.traffic_summary(traffic, "channel")
        _write_table(wb.create_sheet("推广效果"), tr, "渠道推广效果",
                     "ROI=归因GMV/花费；ROI<1 的渠道建议直接控预算",
                     formats={"impressions": "#,##0", "clicks": "#,##0",
                              "spend": "#,##0", "attr_orders": "#,##0",
                              "attr_gmv": "#,##0", "ctr": "0.00%", "cpc": "0.00",
                              "cvr": "0.00%", "roi": "0.00"})

    # 数据质量
    ws9 = wb.create_sheet("数据质量")
    row = _write_table(ws9, ingest_log, "数据采集日志",
                       "每个数据源的接入结果，failed 需当天补齐后重跑")
    if not clean_log.empty:
        _write_table(ws9, clean_log, "数据清洗日志", "命中行数即被规则处理过的记录数",
                     start_row=row)

    # 异常清单
    ws10 = wb.create_sheet("异常清单")
    if issues.empty:
        ws10["A1"] = "本次校验未发现异常"
        ws10["A1"].font = TITLE_FONT
    else:
        show = issues.copy()
        show["status"] = "待处理"
        _write_table(ws10, show, f"异常清单（共 {len(show)} 条）",
                     "处理流程：确认 -> 钉住负责人在系统内标记 acked -> 修复后 fixed",
                     formats={}, color_col="level",
                     colors={"high": "F8CBAD", "medium": "FFE699", "low": "E2EFDA"},
                     max_width=60)
        for r in range(4, 4 + len(show)):
            ws10.row_dimensions[r].height = 30
            for c in (5, 6):
                ws10.cell(row=r, column=c).alignment = Alignment(wrap_text=True,
                                                                vertical="center")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log.info("Excel 报表已生成: %s", path)
    return path


def _html_table(df: pd.DataFrame, percent_cols: tuple[str, ...] = (),
                money_cols: tuple[str, ...] = ()) -> str:
    if df.empty:
        return "<p class='muted'>无数据</p>"
    head = "".join(f"<th>{_cn(c)}</th>" for c in df.columns)
    rows = []
    for r in df.itertuples(index=False):
        cells = []
        for col, v in zip(df.columns, r):
            if isinstance(v, float) and pd.isna(v):
                text = "-"
            elif col in percent_cols:
                text = f"{v:.1%}"
            elif col in money_cols:
                text = f"{v:,.0f}"
            elif isinstance(v, float):
                text = f"{v:,.2f}"
            else:
                text = str(v)
            cells.append(f"<td>{text}</td>")
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return (f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(rows)}</tbody></table>")


def build_html(path: Path, tables: dict[str, pd.DataFrame], issues: pd.DataFrame,
               kpi: dict, meta: dict, cfg: Config, cells: dict[str, Any] | None = None) -> Path:
    orders = tables.get("orders", pd.DataFrame())
    inventory = tables.get("inventory", pd.DataFrame())
    plat = metrics.dimension_summary(orders, "platform", days=30).head(6)
    alert = metrics.inventory_alert(inventory, orders, cfg) if not inventory.empty else pd.DataFrame()
    urgent = alert[alert["status"].isin(["断货", "紧急补货"])].head(8) if not alert.empty else alert

    kpi_cards = "".join(
        f"<div class='card'><div class='label'>{label}</div>"
        f"<div class='value'>{value}</div><div class='sub'>{sub}</div></div>"
        for label, value, sub in [
            ("近30天 GMV", f"¥{kpi.get('gmv', 0):,.0f}",
             f"环比 {kpi.get('gmv_change', 0):+.1%}" if pd.notna(kpi.get("gmv_change")) else "—"),
            ("净销售额", f"¥{kpi.get('net_gmv', 0):,.0f}",
             f"环比 {kpi.get('net_gmv_change', 0):+.1%}" if pd.notna(kpi.get("net_gmv_change")) else "—"),
            ("订单数", f"{kpi.get('orders', 0):,.0f}",
             f"环比 {kpi.get('orders_change', 0):+.1%}" if pd.notna(kpi.get("orders_change")) else "—"),
            ("客单价", f"¥{kpi.get('aov', 0):,.0f}", "实付口径"),
            ("退款率", f"{kpi.get('refund_rate', 0):.2%}", "退款金额/GMV"),
        ])

    issue_html = ""
    if issues.empty:
        issue_html = "<p class='ok'>本次校验未发现异常。</p>"
    else:
        top = issues.head(8)[["level", "rule_name", "message", "suggestion"]].copy()
        top["level"] = top["level"].map({"high": "高", "medium": "中", "low": "低"})
        issue_html = _html_table(top)

    html = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>电商运营日报 {meta.get('date_end', '')}</title>
<style>
  body {{ margin:0; padding:24px; background:#f5f7fa; color:#1f2937;
         font-family:"Microsoft YaHei","PingFang SC",Segoe UI,sans-serif; }}
  .wrap {{ max-width:1080px; margin:0 auto; background:#fff; border-radius:14px;
           box-shadow:0 6px 24px rgba(15,23,42,.08); overflow:hidden; }}
  header {{ background:linear-gradient(120deg,#1f4e79,#2e75b6); color:#fff; padding:24px 28px; }}
  header h1 {{ margin:0 0 6px; font-size:22px; letter-spacing:.5px; }}
  header p {{ margin:0; font-size:12px; opacity:.85; }}
  section {{ padding:18px 28px 6px; }}
  h2 {{ font-size:15px; margin:8px 0 12px; padding-left:9px; border-left:4px solid #c55a11; }}
  .cards {{ display:flex; flex-wrap:wrap; gap:12px; }}
  .card {{ flex:1 1 150px; background:#f8fbff; border:1px solid #ddebf7; border-radius:10px;
           padding:12px 14px; }}
  .card .label {{ font-size:12px; color:#6b7280; }}
  .card .value {{ font-size:20px; font-weight:700; color:#1f4e79; margin:4px 0 2px; }}
  .card .sub {{ font-size:11px; color:#9ca3af; }}
  table {{ width:100%; border-collapse:collapse; font-size:12px; margin-bottom:10px; }}
  th {{ background:#1f4e79; color:#fff; padding:8px; text-align:left; font-weight:600; }}
  td {{ padding:7px 8px; border-bottom:1px solid #eef2f7; }}
  tbody tr:nth-child(even) {{ background:#fafcff; }}
  .muted {{ color:#9ca3af; font-size:12px; }}
  .ok {{ color:#2f7d32; font-size:13px; }}
  footer {{ padding:14px 28px 22px; font-size:11px; color:#9ca3af; }}
</style></head>
<body><div class="wrap">
<header>
  <h1>电商运营数据日报</h1>
  <p>数据区间 {meta.get('date_start', '-')} ~ {meta.get('date_end', '-')} ·
     生成时间 {meta.get('finished_at', '-')} · 批次 {meta.get('run_id', '-')} ·
     数据质量分 {metrics.data_quality_score(issues):.0f}</p>
</header>
<section><h2>核心指标</h2><div class="cards">{kpi_cards}</div></section>
<section><h2>平台表现（近 30 天）</h2>
{_html_table(plat[['platform', 'gmv', 'orders', 'aov', 'margin_rate', 'refund_rate', 'gmv_share']],
             percent_cols=('margin_rate', 'refund_rate', 'gmv_share'),
             money_cols=('gmv', 'orders', 'aov'))}</section>
<section><h2>库存预警（需立即处理）</h2>
{_html_table(urgent[['sku_id', 'warehouse', 'on_hand', 'cover_days', 'status', 'suggestion']]
             if not urgent.empty else urgent)}</section>
<section><h2>异常提醒（Top 8）</h2>{issue_html}</section>
<footer>本报表由 ecomflow-auto 自动生成；口径问题请联系数据负责人。</footer>
</div></body></html>"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(html, encoding="utf-8")
    log.info("HTML 摘要已生成: %s", path)
    return path


def generate(tables: dict[str, pd.DataFrame], issues: pd.DataFrame, kpi: dict,
             clean_log: pd.DataFrame, ingest_log: pd.DataFrame, meta: dict,
             cfg: Config | None = None) -> dict[str, Path]:
    """生成 Excel + HTML 两个交付物，返回路径。"""
    cfg = cfg or get_config()
    out_dir = cfg.path("paths.output")
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    daily = metrics.daily_summary(tables.get("orders", pd.DataFrame()))
    excel_path = out_dir / f"ecom_daily_report_{stamp}.xlsx"
    html_path = out_dir / f"ecom_daily_report_{stamp}.html"
    build_excel(excel_path, tables, issues, daily, kpi, clean_log, ingest_log, meta, cfg)
    build_html(html_path, tables, issues, kpi, meta, cfg)
    return {"excel": excel_path, "html": html_path}
